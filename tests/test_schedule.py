"""Tests for src/schedule.py using a small synthetic workbook with fake names."""
import datetime
from pathlib import Path

import openpyxl
import pytest

from src.masterdata import Flag, MasterData, ScheduleCodeRecord
from src.schedule import parse_schedule_workbook

OUTLETS_CSV = "outlet_id,code_on_clock_file,name_on_record_tab,label_on_schedule,other_names_used,website_listing_crosscheck,status,client_confirm\nAA,AA,Shop AA,AA,,,Per your table,\n"
STAFF_ALIASES_CSV = (
    "id_initials,record_tab,full_name_on_tab,clock_file_names,home_outlet_tab,"
    "schedule_block,evidence,status,client_confirm\n"
    "FAKEONE,FAKEONE,Fake One,FAKEONE (AA),AA,AA,,Matched,\n"
    "FAKETWO,FAKETWO,Fake Two,FAKETWO (AA),AA,AA,,Matched,\n"
)
SCHEDULE_CODES_CSV = (
    "code,count_may_2026,day_type,start,end,hours_in_roster,note\n"
    "AM,5,WORK,10:00,18:00,8.0,ASSUMED\n"
    "OFF,5,OFF,,,,\n"
    "OFF*,1,OFF,,,,asterisk meaning to confirm\n"
)


@pytest.fixture
def master_data(tmp_path):
    from src.masterdata import load_master_data
    (tmp_path / "outlets.csv").write_text(OUTLETS_CSV, encoding="utf-8")
    (tmp_path / "staff_aliases.csv").write_text(STAFF_ALIASES_CSV, encoding="utf-8")
    (tmp_path / "schedule_codes.csv").write_text(SCHEDULE_CODES_CSV, encoding="utf-8")
    return load_master_data(tmp_path)


def _build_synthetic_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active

    ws["A1"] = None
    ws["B1"] = None
    ws["C1"] = "AA"
    ws.merge_cells("C1:D1")

    ws["A2"] = "Date"
    ws["B2"] = "Day"
    ws["C2"] = "FAKEONE"
    ws["E2"] = "FAKETWO"
    ws.merge_cells("E1:F1")
    ws["E1"] = "AA"

    # row 3: BRING FORWARD HRS, no date -> must be skipped
    ws["A3"] = "BRING FORWARD HRS"
    ws["C3"] = 0.0

    # row 4+: data
    ws["A4"] = datetime.datetime(2026, 5, 4)
    ws["B4"] = "Mon"
    ws["C4"] = "AM"
    ws["D4"] = 8.0
    ws["E4"] = "AM"
    ws["F4"] = 8.0

    ws["A5"] = datetime.datetime(2026, 5, 5)
    ws["B5"] = "Tue"
    ws["C5"] = "OFF"
    ws["E5"] = "OFF*"

    ws["A6"] = datetime.datetime(2026, 5, 6)
    ws["B6"] = "Wed"
    ws["C6"] = "12-9PM"  # unresolvable, no asterisk/arrow
    ws["E6"] = "10AM-->"  # unresolvable, arrow

    # weekly total row, no date -> must be skipped
    ws["A7"] = None
    ws["C7"] = 16.0

    wb.save(path)


def test_skips_bring_forward_and_weekly_total_rows(tmp_path, master_data):
    xlsx_path = tmp_path / "schedule.xlsx"
    _build_synthetic_workbook(xlsx_path)

    # "BRING FORWARD HRS" sits at row 3, before FIRST_DATA_ROW (4), so the
    # scan never visits it; only the weekly-total row (7, no date) is
    # counted as skipped during the scan.
    rows, skipped = parse_schedule_workbook(xlsx_path, master_data)
    assert skipped == 1  # the weekly total row
    dates = {r["date"] for r in rows}
    assert dates == {datetime.date(2026, 5, 4), datetime.date(2026, 5, 5), datetime.date(2026, 5, 6)}


def test_person_and_outlet_block_and_hours(tmp_path, master_data):
    xlsx_path = tmp_path / "schedule.xlsx"
    _build_synthetic_workbook(xlsx_path)

    rows, _ = parse_schedule_workbook(xlsx_path, master_data)
    row = next(r for r in rows if r["person"] == "FAKEONE" and r["date"] == datetime.date(2026, 5, 4))
    assert row["outlet_block"] == "AA"
    assert row["raw_code"] == "AM"
    assert row["hours"] == 8.0
    assert isinstance(row["resolved"], ScheduleCodeRecord)
    assert row["resolved"].day_type == "WORK"


def test_known_asterisk_code_resolves(tmp_path, master_data):
    xlsx_path = tmp_path / "schedule.xlsx"
    _build_synthetic_workbook(xlsx_path)

    rows, _ = parse_schedule_workbook(xlsx_path, master_data)
    row = next(r for r in rows if r["person"] == "FAKETWO" and r["date"] == datetime.date(2026, 5, 5))
    assert row["raw_code"] == "OFF*"
    assert isinstance(row["resolved"], ScheduleCodeRecord)
    assert row["resolved"].day_type == "OFF"


def test_unresolvable_codes_become_flags_with_raw_text_kept(tmp_path, master_data):
    xlsx_path = tmp_path / "schedule.xlsx"
    _build_synthetic_workbook(xlsx_path)

    rows, _ = parse_schedule_workbook(xlsx_path, master_data)

    row = next(r for r in rows if r["person"] == "FAKEONE" and r["date"] == datetime.date(2026, 5, 6))
    assert row["raw_code"] == "12-9PM"
    assert isinstance(row["resolved"], Flag)
    assert row["resolved"].kind == "unknown_schedule_code"
    assert "12-9PM" in row["resolved"].detail

    row = next(r for r in rows if r["person"] == "FAKETWO" and r["date"] == datetime.date(2026, 5, 6))
    assert row["raw_code"] == "10AM-->"
    assert isinstance(row["resolved"], Flag)
    assert row["resolved"].kind == "unknown_schedule_code"
    assert "10AM-->" in row["resolved"].detail
