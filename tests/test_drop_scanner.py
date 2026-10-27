"""Tests for src/agents/drop_scanner.py. Synthetic files; fake IDs only."""
import os

import openpyxl
import pytest

import src.daily as daily
from src.agents import drop_scanner


@pytest.fixture
def env(tmp_path, monkeypatch):
    drop, logs = tmp_path / "01_DROP_FOLDER", tmp_path / "07_LOGS"
    (drop / "2026-09").mkdir(parents=True)
    monkeypatch.setattr(daily, "DROP_DIR", drop)
    monkeypatch.setattr(daily, "LOGS_DIR", logs)
    return {"root": tmp_path, "drop": drop, "logs": logs, "m": drop / "2026-09"}


def _roster(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(2, 1, "Date")
    ws.cell(2, 2, "Day")
    wb.save(path)


def _hr(path):
    wb = openpyxl.Workbook()
    wb.active.cell(8, 4, "IN")
    wb.save(path)


def _plain_xlsx(path):
    openpyxl.Workbook().save(path)


def _snapshot(root):
    return {p.relative_to(root).as_posix(): (p.stat().st_size, p.stat().st_mtime_ns)
            for p in root.rglob("*") if p.is_file()}


def _full_month(m):
    _roster(m / "sched.xlsx")
    _hr(m / "hrbook.xlsx")
    (m / "AA-All_Report_x.xls").write_bytes(b"x")


def test_classification_each_class(env):
    m = env["m"]
    _full_month(m)
    _plain_xlsx(m / "other.xlsx")
    _plain_xlsx(m / "q_reply.xlsx")
    (m / "notes.txt").write_text("hi")
    r = drop_scanner.run()
    assert r["status"] == "OK"
    assert r["counts_by_class"] == {"roster": 1, "hr": 1, "punches": 1, "reply": 1, "unknown": 2}
    assert r["scanned"] == r["new"] == 6
    assert any("unknown" in a for a in r["alerts"])
    assert not any("notes" in a or "sched" in a for a in r["alerts"])


def test_idempotent_and_changed_file(env):
    _full_month(env["m"])
    first = drop_scanner.run()
    assert first["status"] == "OK" and first["listing_path"]
    before = sorted(p.name for p in env["logs"].glob("scan_listing_*"))
    second = drop_scanner.run()
    assert second["status"] == "NOTHING_NEW" and second["listing_path"] is None
    assert sorted(p.name for p in env["logs"].glob("scan_listing_*")) == before
    (env["m"] / "AA-All_Report_x.xls").write_bytes(b"xyz")
    third = drop_scanner.run()
    assert third["status"] == "OK" and third["new"] == 1
    assert third["counts_by_class"]["punches"] == 1


def test_mtime_change_detected(env):
    _full_month(env["m"])
    drop_scanner.run()
    f = env["m"] / "AA-All_Report_x.xls"
    st = f.stat()
    os.utime(f, ns=(st.st_atime_ns, st.st_mtime_ns + 5_000_000_000))
    assert drop_scanner.run()["new"] == 1


def test_missing_piece_and_orphan_reply_alerts(env):
    _roster(env["m"] / "sched.xlsx")
    _plain_xlsx(env["drop"] / "loose_reply.xlsx")
    r = drop_scanner.run()
    text = " ".join(r["alerts"])
    assert "2026-09" in text and "hr" in text and "punches" in text
    assert "no matching month" in text
    assert "loose" not in text and "sched" not in text


def test_no_alert_for_complete_month(env):
    _full_month(env["m"])
    assert drop_scanner.run()["alerts"] == []


def test_nothing_copied_and_writes_only_logs(env, capsys):
    _full_month(env["m"])
    before = _snapshot(env["root"])
    drop_before = _snapshot(env["drop"])
    r = drop_scanner.run()
    after = _snapshot(env["root"])
    assert _snapshot(env["drop"]) == drop_before
    added = set(after) - set(before)
    assert added and all(a.startswith("07_LOGS/") for a in added)
    assert all(after[k] == before[k] for k in before)
    assert {"last_scan.json", "drop_scan_log.csv"} <= {a.split("/")[-1] for a in added}
    assert r["listing_path"].startswith(str(env["logs"]))
    out = capsys.readouterr()
    assert out.out == "" and out.err == ""


def test_listing_columns(env):
    _roster(env["m"] / "sched.xlsx")
    r = drop_scanner.run()
    lines = open(r["listing_path"], encoding="utf-8").read().splitlines()
    assert lines[0] == "name,month_folder,size,mtime,class"
    assert lines[1].startswith("sched.xlsx,2026-09,") and lines[1].endswith(",roster")


def test_empty_drop_folder(env):
    r = drop_scanner.run()
    assert r["status"] == "NOTHING_NEW" and r["scanned"] == 0 and r["new"] == 0
    assert not list(env["logs"].glob("scan_listing_*"))


def test_missing_drop_folder_fails_with_log(env, monkeypatch):
    monkeypatch.setattr(daily, "DROP_DIR", env["root"] / "nope")
    r = drop_scanner.run()
    assert r["status"] == "FAILED" and r["alerts"]
    log = (env["logs"] / "drop_scan_log.csv").read_text(encoding="utf-8")
    assert "FAILED" in log


def test_main_returns_zero_on_ok(env):
    _full_month(env["m"])
    assert drop_scanner.main() == 0


def test_main_returns_zero_on_nothing_new(env):
    assert drop_scanner.main() == 0


def test_main_returns_one_on_failure(env, monkeypatch):
    monkeypatch.setattr(daily, "DROP_DIR", env["root"] / "nope")
    assert drop_scanner.main() == 1
