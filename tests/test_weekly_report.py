"""Tests for src/agents/weekly_report.py. Synthetic rows; fake IDs only."""
import datetime
import shutil
from pathlib import Path

import openpyxl
import pytest

import src.daily as daily
from src.agents import weekly_report

D = datetime.date
PROJECT = Path(__file__).resolve().parent.parent


def _row(person, date, flags=()):
    return {"person": person, "date": date, "day": date.strftime("%a"), "schedule_code": "AM",
            "type": "WORK", "shift_start": "10:00", "shift_end": "19:00", "punches": "",
            "clock_outlet": "AA", "system_in": "", "system_out": "", "hq_in": "", "hq_out": "",
            "flag_refs": list(flags)}


def _month_rows(year, month, days, person="FAKEONE", flagged=()):
    return [_row(person, D(year, month, d), [0] if d in flagged else []) for d in days]


@pytest.fixture
def env(tmp_path, monkeypatch):
    drop, reports, logs = tmp_path / "01_DROP_FOLDER", tmp_path / "06_REPORTS", tmp_path / "07_LOGS"
    drop.mkdir()
    monkeypatch.setattr(daily, "DROP_DIR", drop)
    monkeypatch.setattr(daily, "REPORTS_DIR", reports)
    monkeypatch.setattr(daily, "LOGS_DIR", logs)
    import src.run_lock as run_lock
    monkeypatch.setattr(run_lock, "_LOCK_PATH", logs / "engine.lock")
    monkeypatch.setattr(run_lock, "_LOGS_DIR", logs)
    return {"root": tmp_path, "drop": drop, "reports": reports, "logs": logs}


def _fake_engine(monkeypatch, per_month, calls=None):
    def fake(month, force=False, _test_hook=None):
        if calls is not None:
            calls.append((month, force))
        return per_month[month]
    monkeypatch.setattr(daily, "run_month", fake)


def _tree(root: Path):
    return {p.relative_to(root) for p in root.rglob("*") if p.is_file()}


def test_window_filter_keeps_only_the_week(env, monkeypatch):
    (env["drop"] / "2026-10").mkdir()
    # Mon 5 Oct .. Sun 11 Oct is the previous week of Sat 17 Oct.
    rows = _month_rows(2026, 10, range(1, 20), flagged={6, 15})
    calls = []
    _fake_engine(monkeypatch, {"2026-10": {"status": "OK", "day_check_rows": rows, "missing": []}}, calls)
    res = weekly_report.run(as_of=D(2026, 10, 17))
    assert res["status"] == "OK"
    assert calls == [("2026-10", True)]
    assert res["counts"]["person_days"] == 7
    assert res["counts"]["flagged_person_days"] == 1  # day 15 is outside the window
    assert res["counts"]["flags"] == 1
    assert res["output_path"].endswith(str(Path("06_REPORTS") / "2026-10" / "Weekly_2026-W41.xlsx"))
    wb = openpyxl.load_workbook(res["output_path"])
    assert wb.sheetnames == ["Summary", "By person", "By day", "Day check"]
    assert wb["Day check"].max_row == 8  # header + 7 days


def test_cross_month_week_merges_both_months(env, monkeypatch):
    (env["drop"] / "2026-09").mkdir()
    (env["drop"] / "2026-10").mkdir()
    # Previous week of Thu 8 Oct is Mon 28 Sep .. Sun 4 Oct.
    _fake_engine(monkeypatch, {
        "2026-09": {"status": "OK", "missing": [],
                    "day_check_rows": _month_rows(2026, 9, range(25, 31), flagged={29})},
        "2026-10": {"status": "OK", "missing": [],
                    "day_check_rows": _month_rows(2026, 10, range(1, 8), flagged={2})},
    })
    res = weekly_report.run(as_of=D(2026, 10, 8))
    assert res["status"] == "OK"
    assert (res["start"], res["end"]) == ("2026-09-28", "2026-10-04")
    assert res["counts"]["person_days"] == 7  # 28,29,30 Sep + 1,2,3,4 Oct
    assert res["counts"]["flags"] == 2
    # Output goes to the month folder of the END date.
    assert Path(res["output_path"]).parent.name == "2026-10"


def test_missing_month_folder_gives_partial_with_covered_dates(env, monkeypatch):
    (env["drop"] / "2026-10").mkdir()  # 2026-09 is missing
    _fake_engine(monkeypatch, {"2026-10": {"status": "OK", "missing": [],
                                           "day_check_rows": _month_rows(2026, 10, range(1, 8))}})
    res = weekly_report.run(as_of=D(2026, 10, 8))
    assert res["status"] == "PARTIAL"
    assert res["missing_months"] == ["2026-09"]
    assert (res["covered_start"], res["covered_end"]) == ("2026-10-01", "2026-10-04")
    assert "_DRAFT" in res["output_path"]
    wb = openpyxl.load_workbook(res["output_path"])
    text = " ".join(str(c.value) for row in wb["Summary"].iter_rows() for c in row if c.value)
    assert "2026-10-01 to 2026-10-04" in text and "2026-09" in text


def test_no_month_folder_at_all_writes_nothing(env, monkeypatch):
    _fake_engine(monkeypatch, {})
    res = weekly_report.run(as_of=D(2026, 10, 17))
    assert res["status"] == "FAILED"
    assert res["output_path"] is None
    assert not env["reports"].exists()
    assert (env["logs"] / "weekly_log.csv").exists()


def test_engine_failed_logs_and_writes_no_weekly_file(env, monkeypatch):
    (env["drop"] / "2026-10").mkdir()
    _fake_engine(monkeypatch, {"2026-10": {"status": "FAILED", "error": "boom", "missing": []}})
    res = weekly_report.run(as_of=D(2026, 10, 17))
    assert res["status"] == "FAILED" and "boom" in res["error"]
    assert not env["reports"].exists() or not list(env["reports"].rglob("Weekly_*"))
    log = (env["logs"] / "weekly_log.csv").read_text(encoding="utf-8")
    assert "FAILED" in log and "boom" in log


def test_rerun_does_not_overwrite(env, monkeypatch):
    (env["drop"] / "2026-10").mkdir()
    _fake_engine(monkeypatch, {"2026-10": {"status": "OK", "missing": [],
                                           "day_check_rows": _month_rows(2026, 10, range(5, 12))}})
    first = weekly_report.run(as_of=D(2026, 10, 17))
    path = Path(first["output_path"])
    before = (path.stat().st_mtime_ns, path.read_bytes())
    second = weekly_report.run(as_of=D(2026, 10, 17))
    assert second["status"] == "SKIPPED_EXISTS"
    assert (path.stat().st_mtime_ns, path.read_bytes()) == before


def test_draft_replaced_by_final(env, monkeypatch):
    (env["drop"] / "2026-10").mkdir()
    _fake_engine(monkeypatch, {
        "2026-10": {"status": "OK", "missing": [],
                    "day_check_rows": _month_rows(2026, 10, range(1, 8))},
    })
    partial = weekly_report.run(as_of=D(2026, 10, 8))
    assert partial["status"] == "PARTIAL"
    draft = Path(partial["output_path"])
    assert draft.exists() and "_DRAFT" in draft.name
    (env["drop"] / "2026-09").mkdir()
    _fake_engine(monkeypatch, {
        "2026-09": {"status": "OK", "missing": [],
                    "day_check_rows": _month_rows(2026, 9, range(25, 31))},
        "2026-10": {"status": "OK", "missing": [],
                    "day_check_rows": _month_rows(2026, 10, range(1, 8))},
    })
    final = weekly_report.run(as_of=D(2026, 10, 8))
    assert final["status"] == "OK"
    assert "_DRAFT" not in final["output_path"]
    assert not draft.exists()
    assert Path(final["output_path"]).exists()


def test_writes_only_under_reports_and_logs(env, monkeypatch):
    (env["drop"] / "2026-10").mkdir()
    (env["drop"] / "2026-10" / "keep.txt").write_text("x", encoding="utf-8")
    before = _tree(env["root"])
    _fake_engine(monkeypatch, {"2026-10": {"status": "OK", "missing": [],
                                           "day_check_rows": _month_rows(2026, 10, range(5, 12))}})
    weekly_report.run(as_of=D(2026, 10, 17))
    new = _tree(env["root"]) - before
    assert new and all(p.parts[0] in ("06_REPORTS", "07_LOGS") for p in new)
    assert (env["drop"] / "2026-10" / "keep.txt").read_text(encoding="utf-8") == "x"


def test_run_prints_nothing(env, monkeypatch, capsys):
    (env["drop"] / "2026-10").mkdir()
    _fake_engine(monkeypatch, {"2026-10": {"status": "OK", "missing": [], "day_check_rows": []}})
    weekly_report.run(as_of=D(2026, 10, 17))
    out = capsys.readouterr()
    assert out.out == "" and out.err == ""


def test_may_sample_real_engine_in_temp_folder(env, monkeypatch):
    sample = PROJECT / "01_DROP_FOLDER" / "2026-05"
    master = PROJECT / "03_MASTER_DATA"
    if not sample.is_dir() or not master.is_dir() or not list(sample.glob("*.xls")):
        pytest.skip("May 2026 sample not present")
    shutil.copytree(sample, env["drop"] / "2026-05")
    tmp_master = env["root"] / "03_MASTER_DATA"
    shutil.copytree(master, tmp_master)
    monkeypatch.setattr(daily, "MASTER_DATA_DIR", tmp_master)
    monkeypatch.setattr(daily, "MANIFEST_PATH", env["logs"] / "processed_manifest.json")
    monkeypatch.setattr(daily, "RUN_LOG_PATH", env["logs"] / "run_log.csv")
    monkeypatch.setattr(daily, "DAILY_TASK_LOG_PATH", env["logs"] / "daily_task.log")
    import src.run_lock as run_lock
    monkeypatch.setattr(run_lock, "_LOCK_PATH", env["logs"] / "engine.lock")
    monkeypatch.setattr(run_lock, "_LOGS_DIR", env["logs"])
    monkeypatch.setattr(daily, "APPROVALS_PATH", tmp_master / "approvals.csv")
    res = weekly_report.run(as_of=D(2026, 5, 14))  # previous week: 4-10 May
    assert res["status"] in ("OK", "PARTIAL")
    assert res["counts"]["person_days"] > 0
    assert Path(res["output_path"]).name == "Weekly_2026-W19.xlsx"


def test_main_returns_zero_on_ok(env, monkeypatch):
    (env["drop"] / "2026-10").mkdir()
    _fake_engine(monkeypatch, {"2026-10": {"status": "OK", "missing": [],
                                           "day_check_rows": _month_rows(2026, 10, range(5, 12))}})
    assert weekly_report.main() == 0


def test_main_returns_one_on_failure(env, monkeypatch):
    _fake_engine(monkeypatch, {})
    assert weekly_report.main() == 1
