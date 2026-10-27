"""Tests for src/agents/monthly_report.py. Synthetic rows; fake IDs only."""
import datetime
import shutil
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest

import src.daily as daily
from src.agents import monthly_report

D = datetime.date
PROJECT = Path(__file__).resolve().parent.parent


def _row(person, date, flags=(), outlet="AA"):
    return {"person": person, "date": date, "day": date.strftime("%a"), "schedule_code": "AM",
            "type": "WORK", "shift_start": "10:00", "shift_end": "19:00", "punches": "",
            "clock_outlet": outlet, "system_in": "", "system_out": "", "hq_in": "", "hq_out": "",
            "flag_refs": list(flags)}


def _month_rows(year, month, days, person="FAKEONE", flagged=(), outlet="AA"):
    return [_row(person, D(year, month, d), [0] if d in flagged else [], outlet) for d in days]


@pytest.fixture
def env(tmp_path, monkeypatch):
    drop, reports, logs = tmp_path / "01_DROP_FOLDER", tmp_path / "06_REPORTS", tmp_path / "07_LOGS"
    drop.mkdir()
    monkeypatch.setattr(daily, "DROP_DIR", drop)
    monkeypatch.setattr(daily, "REPORTS_DIR", reports)
    monkeypatch.setattr(daily, "LOGS_DIR", logs)
    import src.run_lock as run_lock
    monkeypatch.setattr(run_lock, "_LOCK_PATH", logs / "engine.lock")
    monkeypatch.setattr(run_lock, "_LOGS_DIR", logs)
    return {"root": tmp_path, "drop": drop, "reports": reports, "logs": logs}


def _fake_engine(monkeypatch, per_month, calls=None):
    def fake(month, force=False, _test_hook=None):
        if calls is not None:
            calls.append((month, force))
        return per_month[month]
    monkeypatch.setattr(daily, "run_month", fake)


def _tree(root: Path):
    return {p.relative_to(root) for p in root.rglob("*") if p.is_file()}


def test_previous_month_counts_and_sheets(env, monkeypatch):
    (env["drop"] / "2026-09").mkdir()
    rows = (_month_rows(2026, 9, range(1, 11), flagged={3, 4}, outlet="AA")
            + _month_rows(2026, 9, range(1, 6), person="FAKETWO", outlet="BB"))
    calls = []
    _fake_engine(monkeypatch, {"2026-09": {"status": "OK", "day_check_rows": rows, "missing": []}}, calls)
    res = monthly_report.run(as_of=D(2026, 10, 18))
    assert res["status"] == "OK"
    assert calls == [("2026-09", True)]
    assert res["counts"] == {"person_days": 15, "flagged_person_days": 2, "flags": 2, "people": 2}
    assert res["month_flags"] is None
    assert res["output_path"].endswith(str(Path("06_REPORTS") / "2026-09" / "Monthly_2026-09.xlsx"))
    wb = openpyxl.load_workbook(res["output_path"])
    assert wb.sheetnames == ["Summary", "By person", "By day", "By outlet", "Day check"]
    assert wb["Day check"].max_row == 16
    outlets = {r[0].value: r[1].value for r in wb["By outlet"].iter_rows(min_row=2)}
    assert outlets == {"AA": 10, "BB": 5}
    log = (env["logs"] / "monthly_log.csv").read_text(encoding="utf-8")
    assert log.startswith("run_time,window,status,detail") and "2026-09_Monthly,OK" in log


def test_december_to_january_year_boundary(env, monkeypatch):
    (env["drop"] / "2025-12").mkdir()
    calls = []
    _fake_engine(monkeypatch, {"2025-12": {"status": "OK", "missing": [],
                                           "day_check_rows": _month_rows(2025, 12, range(1, 32))}}, calls)
    res = monthly_report.run(as_of=D(2026, 1, 5))
    assert calls == [("2025-12", True)]
    assert (res["start"], res["end"]) == ("2025-12-01", "2025-12-31")
    assert res["counts"]["person_days"] == 31
    assert Path(res["output_path"]).name == "Monthly_2025-12.xlsx"
    assert Path(res["output_path"]).parent.name == "2025-12"


def test_month_level_flags_included_when_present(env, monkeypatch):
    (env["drop"] / "2026-09").mkdir()
    flags = [SimpleNamespace(severity="Check", remark_type="Possible error"),
             SimpleNamespace(severity="Check", remark_type="Pending question"),
             SimpleNamespace(severity="Info", remark_type="Confirmed rule")]
    _fake_engine(monkeypatch, {"2026-09": {"status": "OK", "missing": [], "all_flags": flags,
                                           "day_check_rows": _month_rows(2026, 9, range(1, 4))}})
    res = monthly_report.run(as_of=D(2026, 10, 18))
    assert res["month_flags"] == {"red": 1, "amber": 1, "green": 0, "info": 1}
    wb = openpyxl.load_workbook(res["output_path"])
    text = " ".join(str(c.value) for row in wb["Summary"].iter_rows() for c in row if c.value)
    assert "Month-level flags" in text


def test_bad_all_flags_is_ignored(env, monkeypatch):
    (env["drop"] / "2026-09").mkdir()
    _fake_engine(monkeypatch, {"2026-09": {"status": "OK", "missing": [], "all_flags": [object()],
                                           "day_check_rows": _month_rows(2026, 9, range(1, 4))}})
    res = monthly_report.run(as_of=D(2026, 10, 18))
    assert res["status"] == "OK" and res["month_flags"] is None


def test_engine_partial_gives_partial(env, monkeypatch):
    (env["drop"] / "2026-09").mkdir()
    _fake_engine(monkeypatch, {"2026-09": {"status": "PARTIAL", "missing": ["no schedule"],
                                           "day_check_rows": _month_rows(2026, 9, range(1, 4))}})
    res = monthly_report.run(as_of=D(2026, 10, 18))
    assert res["status"] == "PARTIAL" and res["missing"] == ["no schedule"]
    assert "_DRAFT" in res["output_path"]


def test_missing_month_folder_writes_nothing(env, monkeypatch):
    _fake_engine(monkeypatch, {})
    res = monthly_report.run(as_of=D(2026, 10, 18))
    assert res["status"] == "FAILED" and res["output_path"] is None
    assert not env["reports"].exists()
    assert (env["logs"] / "monthly_log.csv").exists()


def test_engine_failed_logs_and_writes_no_file(env, monkeypatch):
    (env["drop"] / "2026-09").mkdir()
    _fake_engine(monkeypatch, {"2026-09": {"status": "FAILED", "error": "boom", "missing": []}})
    res = monthly_report.run(as_of=D(2026, 10, 18))
    assert res["status"] == "FAILED" and "boom" in res["error"]
    assert not env["reports"].exists() or not list(env["reports"].rglob("Monthly_*"))
    log = (env["logs"] / "monthly_log.csv").read_text(encoding="utf-8")
    assert "FAILED" in log and "boom" in log


def test_rerun_does_not_overwrite(env, monkeypatch):
    (env["drop"] / "2026-09").mkdir()
    _fake_engine(monkeypatch, {"2026-09": {"status": "OK", "missing": [],
                                           "day_check_rows": _month_rows(2026, 9, range(1, 8))}})
    first = monthly_report.run(as_of=D(2026, 10, 18))
    path = Path(first["output_path"])
    before = (path.stat().st_mtime_ns, path.read_bytes())
    second = monthly_report.run(as_of=D(2026, 10, 18))
    assert second["status"] == "SKIPPED_EXISTS"
    assert (path.stat().st_mtime_ns, path.read_bytes()) == before


def test_draft_replaced_by_final_on_month_end(env, monkeypatch):
    (env["drop"] / "2026-09").mkdir()
    rows_partial = _month_rows(2026, 9, range(1, 4))
    _fake_engine(monkeypatch, {"2026-09": {"status": "PARTIAL", "missing": ["no schedule"],
                                           "day_check_rows": rows_partial}})
    partial = monthly_report.run(as_of=D(2026, 10, 18))
    assert partial["status"] == "PARTIAL"
    draft = Path(partial["output_path"])
    assert draft.exists() and "_DRAFT" in draft.name
    rows_full = _month_rows(2026, 9, range(1, 31))
    _fake_engine(monkeypatch, {"2026-09": {"status": "OK", "missing": [],
                                           "day_check_rows": rows_full}})
    final = monthly_report.run(as_of=D(2026, 10, 18))
    assert final["status"] == "OK"
    assert "_DRAFT" not in final["output_path"]
    assert not draft.exists()
    assert Path(final["output_path"]).exists()


def test_writes_only_under_reports_and_logs(env, monkeypatch):
    (env["drop"] / "2026-09").mkdir()
    (env["drop"] / "2026-09" / "keep.txt").write_text("x", encoding="utf-8")
    before = _tree(env["root"])
    _fake_engine(monkeypatch, {"2026-09": {"status": "OK", "missing": [],
                                           "day_check_rows": _month_rows(2026, 9, range(1, 8))}})
    monthly_report.run(as_of=D(2026, 10, 18))
    new = _tree(env["root"]) - before
    assert new and all(p.parts[0] in ("06_REPORTS", "07_LOGS") for p in new)
    assert (env["drop"] / "2026-09" / "keep.txt").read_text(encoding="utf-8") == "x"


def test_run_prints_nothing(env, monkeypatch, capsys):
    (env["drop"] / "2026-09").mkdir()
    _fake_engine(monkeypatch, {"2026-09": {"status": "OK", "missing": [], "day_check_rows": []}})
    monthly_report.run(as_of=D(2026, 10, 18))
    out = capsys.readouterr()
    assert out.out == "" and out.err == ""


def test_may_sample_real_engine_in_temp_folder(env, monkeypatch):
    sample = PROJECT / "01_DROP_FOLDER" / "2026-05"
    master = PROJECT / "03_MASTER_DATA"
    if not sample.is_dir() or not master.is_dir() or not list(sample.glob("*.xls")):
        pytest.skip("May 2026 sample not present")
    shutil.copytree(sample, env["drop"] / "2026-05")
    tmp_master = env["root"] / "03_MASTER_DATA"
    shutil.copytree(master, tmp_master)
    monkeypatch.setattr(daily, "MASTER_DATA_DIR", tmp_master)
    monkeypatch.setattr(daily, "MANIFEST_PATH", env["logs"] / "processed_manifest.json")
    monkeypatch.setattr(daily, "RUN_LOG_PATH", env["logs"] / "run_log.csv")
    monkeypatch.setattr(daily, "DAILY_TASK_LOG_PATH", env["logs"] / "daily_task.log")
    import src.run_lock as run_lock
    monkeypatch.setattr(run_lock, "_LOCK_PATH", env["logs"] / "engine.lock")
    monkeypatch.setattr(run_lock, "_LOGS_DIR", env["logs"])
    monkeypatch.setattr(daily, "APPROVALS_PATH", tmp_master / "approvals.csv")
    res = monthly_report.run(as_of=D(2026, 6, 10))  # previous month: May 2026
    assert res["status"] in ("OK", "PARTIAL")
    assert res["counts"]["person_days"] > 0
    assert Path(res["output_path"]).name == "Monthly_2026-05.xlsx"


def test_main_returns_zero_on_ok(env, monkeypatch):
    (env["drop"] / "2026-09").mkdir()
    _fake_engine(monkeypatch, {"2026-09": {"status": "OK", "missing": [],
                                           "day_check_rows": _month_rows(2026, 9, range(1, 8))}})
    assert monthly_report.main() == 0


def test_main_returns_one_on_failure(env, monkeypatch):
    _fake_engine(monkeypatch, {})
    assert monthly_report.main() == 1
