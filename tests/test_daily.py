"""Tests for src/daily.py using synthetic data with fake names only."""
import datetime
import json
import time
from pathlib import Path

import openpyxl
import pytest

import src.daily as daily
from src.dayengine import DayFlag

OUTLETS_CSV = (
    "outlet_id,code_on_clock_file,name_on_record_tab,label_on_schedule,other_names_used,"
    "website_listing_crosscheck,status,client_confirm\n"
    "AA,AA,Shop AA,AA,,,Per your table,\n"
)
STAFF_ALIASES_CSV = (
    "id_initials,record_tab,full_name_on_tab,clock_file_names,home_outlet_tab,"
    "schedule_block,evidence,status,client_confirm\n"
    "FAKEONE,FAKEONE,Fake One,(no clock file in batch),AA,AA,,Draft,\n"
)
SCHEDULE_CODES_CSV = (
    "code,count_may_2026,day_type,start,end,hours_in_roster,note\n"
    "UPL,5,LEAVE,,,8.0,Q11\n"
)
RULES_TOML_TEMPLATE = """
[meta]
schema_version = 1
allowed_status = ["CONFIRMED", "ASSUMED", "OPEN"]

[rules.morning_open]
label = "x"
value = "10:00"
status = "CONFIRMED"
question = "C6"
source = "fake"

[rules.early_clock_in_counts_as_open]
label = "x"
value = true
status = "CONFIRMED"
question = "C6"
source = "fake"

[rules.late_grace_minutes]
label = "x"
value = 15
status = "CONFIRMED"
question = "C5"
source = "fake"

[rules.late_clears_if_approved]
label = "x"
value = true
status = "CONFIRMED"
question = "C5"
source = "fake"

[rules.transfer_is_error]
label = "x"
value = false
status = "CONFIRMED"
question = "C4"
source = "fake"

[rules.ot_sum_from_clock]
label = "x"
value = true
status = "CONFIRMED"
question = "C7"
source = "fake"

[rules.late_deduction_mode]
label = "x"
value = "show_both"
options = ["actual_minutes", "round_up", "show_both"]
status = "OPEN"
question = "Q1"
source = ""

[rules.late_minutes_counted_from]
label = "x"
value = "opening"
options = ["opening", "end_of_grace"]
status = "ASSUMED"
question = "Q1"
source = ""

[rules.grace_applies_to_all_shifts]
label = "x"
value = true
options = [true, false]
status = "ASSUMED"
question = "Q1"
source = ""

[rules.late_round_up_to_minutes]
label = "x"
value = 30
status = "OPEN"
question = "Q1"
source = ""

[rules.approval_source]
label = "x"
value = "none"
options = ["none", "approvals_csv"]
status = "OPEN"
question = "Q8"
source = ""

[rules.count_time_after_end]
label = "x"
value = false
status = "ASSUMED"
question = "Q2"
source = ""

[rules.ot_rate_day_types]
label = "x"
value = ""
status = "OPEN"
question = "Q3"
source = ""

[rules.ot_reference_only]
label = "x"
value = true
status = "CONFIRMED"
question = "C7"
source = "fake"

[rules.single_punch_classification]
label = "x"
value = "closer_to_boundary"
status = "ASSUMED"
question = "Q12"
source = ""

[rules.single_punch_tie_minutes]
label = "x"
value = 60
status = "ASSUMED"
question = "Q12"
source = ""

[rules.duplicate_punch_window_minutes]
label = "x"
value = 5
status = "ASSUMED"
question = "Q15"
source = ""

[rules.duplicate_punch_keep]
label = "x"
value = "first"
options = ["first", "last"]
status = "ASSUMED"
question = "Q15"
source = ""

[rules.early_clock_in_counts_as_start_all_shifts]
label = "x"
value = true
status = "ASSUMED"
question = "Q1b"
source = ""

[rules.standard_day_hours]
label = "x"
value = 8.5
status = "{standard_day_hours_status}"
question = "Q11"
source = "{standard_day_hours_source}"

[rules.leave_day_hours]
label = "x"
value = {leave_day_hours}
status = "{leave_day_hours_status}"
question = "Q11"
source = "{leave_day_hours_source}"

[rules.upl_absent_unpaid]
label = "x"
value = "unknown"
options = ["yes", "no", "unknown"]
status = "OPEN"
question = "Q11"
source = ""

[rules.excluded_staff_ids]
label = "x"
value = []
status = "ASSUMED"
question = "Q6"
source = ""
"""


def _write_master_data(root: Path):
    root.mkdir(parents=True, exist_ok=True)
    (root / "outlets.csv").write_text(OUTLETS_CSV, encoding="utf-8")
    (root / "staff_aliases.csv").write_text(STAFF_ALIASES_CSV, encoding="utf-8")
    (root / "schedule_codes.csv").write_text(SCHEDULE_CODES_CSV, encoding="utf-8")


def _write_rules(
    root: Path, leave_day_hours=8.5, leave_day_hours_status="ASSUMED", leave_day_hours_source="",
    standard_day_hours_status="ASSUMED", standard_day_hours_source="",
):
    root.mkdir(parents=True, exist_ok=True)
    (root / "rules.toml").write_text(
        RULES_TOML_TEMPLATE.format(
            leave_day_hours=leave_day_hours,
            leave_day_hours_status=leave_day_hours_status,
            leave_day_hours_source=leave_day_hours_source,
            standard_day_hours_status=standard_day_hours_status,
            standard_day_hours_source=standard_day_hours_source,
        ),
        encoding="utf-8",
    )


def _write_schedule_workbook(path: Path, month: datetime.datetime, day: int = 4):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["C1"] = "AA"
    ws.merge_cells("C1:D1")
    ws["A2"] = "Date"
    ws["B2"] = "Day"
    ws["C2"] = "FAKEONE"
    date = datetime.datetime(month.year, month.month, day)
    ws["A4"] = date
    ws["B4"] = date.strftime("%a")
    ws["C4"] = "UPL"
    wb.save(path)


def _write_hr_workbook(path: Path, month: datetime.datetime, day: int = 4, working_hours_typed_hours: float = 8.0):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FAKEONE"
    ws.cell(2, 4, month)
    ws.cell(3, 4, "Fake One")
    ws.cell(3, 8, "Shop AA")
    ws.cell(8, 4, "IN")
    ws.cell(8, 5, "OUT")
    ws.cell(8, 6, "Working")
    ws.cell(8, 8, "OT 1.5")
    row = 9 + (day - 1)  # FIRST_DATA_ROW is 9; day = row - 8
    ws.cell(row, 3, day)
    hh = int(working_hours_typed_hours)
    mm = int(round((working_hours_typed_hours - hh) * 60))
    ws.cell(row, 6, datetime.time(hh, mm))  # typed, not a formula
    ws.cell(row, 7, 1.5)
    wb.save(path)


@pytest.fixture(autouse=True)
def _isolate_daily_paths(tmp_path, monkeypatch):
    drop = tmp_path / "01_DROP_FOLDER"
    master = tmp_path / "03_MASTER_DATA"
    reports = tmp_path / "06_REPORTS"
    logs = tmp_path / "07_LOGS"
    drop.mkdir()
    _write_master_data(master)
    _write_rules(master)

    monkeypatch.setattr(daily, "DROP_DIR", drop)
    monkeypatch.setattr(daily, "MASTER_DATA_DIR", master)
    monkeypatch.setattr(daily, "REPORTS_DIR", reports)
    monkeypatch.setattr(daily, "LOGS_DIR", logs)
    monkeypatch.setattr(daily, "MANIFEST_PATH", logs / "processed_manifest.json")
    monkeypatch.setattr(daily, "RUN_LOG_PATH", logs / "run_log.csv")
    monkeypatch.setattr(daily, "DAILY_TASK_LOG_PATH", logs / "daily_task.log")
    import src.run_lock as run_lock
    monkeypatch.setattr(run_lock, "_LOCK_PATH", logs / "engine.lock")
    monkeypatch.setattr(run_lock, "_LOGS_DIR", logs)
    monkeypatch.setattr(daily, "APPROVALS_PATH", master / "approvals.csv")
    return {"drop": drop, "master": master, "reports": reports, "logs": logs}


def _make_full_month(drop: Path, month_str: str):
    month_dir = drop / month_str
    month_dir.mkdir(parents=True, exist_ok=True)
    month = datetime.datetime.strptime(month_str, "%Y-%m")
    _write_schedule_workbook(month_dir / "schedule.xlsx", month)
    _write_hr_workbook(month_dir / "hr.xlsx", month)
    return month_dir


def test_new_month_runs_and_writes_outputs(_isolate_daily_paths):
    paths = _isolate_daily_paths
    _make_full_month(paths["drop"], "2026-05")
    result = daily.process_month("2026-05", force=False)
    assert result["status"] in ("OK", "PARTIAL")
    out_dir = paths["reports"] / "2026-05"
    assert (out_dir / "Attendance_Output_Report_2026-05.xlsx").exists()
    assert (out_dir / "AGGREGATE_SUMMARY.json").exists()
    assert (out_dir / "DAILY_NOTE_basic.txt").exists()
    assert (paths["reports"] / "LATEST" / "AGGREGATE_SUMMARY.json").exists()


def test_nothing_new_on_second_run(_isolate_daily_paths):
    paths = _isolate_daily_paths
    _make_full_month(paths["drop"], "2026-05")
    daily.process_month("2026-05", force=False)
    result = daily.process_month("2026-05", force=False)
    assert result["status"] == "NOTHING_NEW"


def test_missing_hr_workbook_gives_partial_with_missing_sentence(_isolate_daily_paths):
    paths = _isolate_daily_paths
    month_dir = paths["drop"] / "2026-05"
    month_dir.mkdir(parents=True)
    _write_schedule_workbook(month_dir / "schedule.xlsx", datetime.datetime(2026, 5, 1))
    result = daily.process_month("2026-05", force=False)
    assert result["status"] == "PARTIAL"
    assert any("HR record workbook" in m for m in result["missing"])


def test_unknown_file_goes_to_needs_attention(_isolate_daily_paths):
    paths = _isolate_daily_paths
    month_dir = _make_full_month(paths["drop"], "2026-05")
    (month_dir / "mystery.xlsx").write_text("not a real workbook", encoding="utf-8")
    import openpyxl as _opx
    wb = _opx.Workbook()
    wb.active["A1"] = "nothing useful here"
    wb.save(month_dir / "mystery.xlsx")
    daily.process_month("2026-05", force=False)
    needs_attention = (paths["reports"] / "2026-05" / "NEEDS_ATTENTION.txt").read_text(encoding="utf-8")
    assert "mystery.xlsx" in needs_attention
    assert not (month_dir / "NEEDS_ATTENTION.txt").exists()


def test_input_changed_during_run_fails(_isolate_daily_paths):
    paths = _isolate_daily_paths
    month_dir = _make_full_month(paths["drop"], "2026-05")
    hr_path = month_dir / "hr.xlsx"

    def _mutate(signatures_before, tracked_paths):
        time.sleep(0.01)
        with open(hr_path, "ab") as f:
            f.write(b"x")

    result = daily.run_month("2026-05", force=False, _test_hook=_mutate)
    assert result["status"] == "FAILED"
    assert "changed" in result["error"].lower()


def test_aggregate_summary_has_no_staff_names_or_ids(_isolate_daily_paths):
    paths = _isolate_daily_paths
    _make_full_month(paths["drop"], "2026-05")
    daily.process_month("2026-05", force=False)
    summary_text = (paths["reports"] / "2026-05" / "AGGREGATE_SUMMARY.json").read_text(encoding="utf-8")
    note_text = (paths["reports"] / "2026-05" / "DAILY_NOTE_basic.txt").read_text(encoding="utf-8")
    assert "FAKEONE" not in summary_text
    assert "Fake One" not in summary_text
    assert "FAKEONE" not in note_text
    assert "Fake One" not in note_text


def test_changing_leave_day_hours_changes_amber_green_split_with_no_code_change(_isolate_daily_paths):
    paths = _isolate_daily_paths
    _make_full_month(paths["drop"], "2026-05")

    _write_rules(paths["master"], leave_day_hours=8.5, leave_day_hours_status="ASSUMED")
    result_assumed = daily.run_month("2026-05", force=True)
    counts_assumed, _ = daily._counts_by_severity_and_remark(result_assumed["all_flags"])

    _write_rules(
        paths["master"], leave_day_hours=8.5,
        leave_day_hours_status="CONFIRMED", leave_day_hours_source="Client replied",
        standard_day_hours_status="CONFIRMED", standard_day_hours_source="Client replied",
    )
    result_confirmed = daily.run_month("2026-05", force=True)
    counts_confirmed, _ = daily._counts_by_severity_and_remark(result_confirmed["all_flags"])

    assert counts_assumed["amber"] > counts_confirmed["amber"]
    assert counts_confirmed["green"] > counts_assumed["green"]


def test_duplicate_file_names_give_partial_and_needs_attention(_isolate_daily_paths, monkeypatch):
    paths = _isolate_daily_paths
    month_dir = _make_full_month(paths["drop"], "2026-05")
    schedule_path = month_dir / "schedule.xlsx"
    hr_path = month_dir / "hr.xlsx"
    # simulate a collision: a clock file sharing hr.xlsx's basename, from
    # elsewhere, so the manifest's name-only key would otherwise mask it
    other_dir = paths["drop"] / "elsewhere"
    other_dir.mkdir(parents=True)
    duplicate = other_dir / "hr.xlsx"
    duplicate.write_bytes(hr_path.read_bytes())

    monkeypatch.setattr(
        daily, "_classify_month_folder",
        lambda md: ([duplicate], schedule_path, hr_path, []),
    )
    # the duplicate is a copy of the HR workbook's bytes under a clock
    # file's name, which capture_folder cannot parse as a clock export;
    # that is not what this test is about, so bypass real clock parsing.
    monkeypatch.setattr(daily, "capture_folder", lambda folder, pattern="*": ([], []))
    result = daily.process_month("2026-05", force=True)
    assert result["status"] == "PARTIAL"
    assert any("duplicate file name" in m.lower() for m in result["missing"])
    needs_attention = (paths["reports"] / "2026-05" / "NEEDS_ATTENTION.txt").read_text(encoding="utf-8")
    assert "hr.xlsx" in needs_attention
    assert not (month_dir / "NEEDS_ATTENTION.txt").exists()


def test_nothing_new_repeats_prior_partial_and_missing(_isolate_daily_paths):
    paths = _isolate_daily_paths
    month_dir = paths["drop"] / "2026-05"
    month_dir.mkdir(parents=True)
    _write_schedule_workbook(month_dir / "schedule.xlsx", datetime.datetime(2026, 5, 1))
    # no HR workbook -> first run is PARTIAL with a missing-HR sentence
    first = daily.process_month("2026-05", force=False)
    assert first["status"] == "PARTIAL"
    assert any("HR record workbook" in m for m in first["missing"])

    second = daily.process_month("2026-05", force=False)
    assert second["status"] == "NOTHING_NEW"
    assert any("HR record workbook" in m for m in second["missing"])

    summary_text = (paths["reports"] / "2026-05" / "AGGREGATE_SUMMARY.json").read_text(encoding="utf-8")
    assert "HR record workbook" in summary_text


def test_month_mismatch_gives_partial(_isolate_daily_paths):
    paths = _isolate_daily_paths
    month_dir = paths["drop"] / "2026-05"
    month_dir.mkdir(parents=True)
    # schedule and HR workbook both say a DIFFERENT month than the folder
    _write_schedule_workbook(month_dir / "schedule.xlsx", datetime.datetime(2026, 4, 1))
    _write_hr_workbook(month_dir / "hr.xlsx", datetime.datetime(2026, 4, 1))
    result = daily.process_month("2026-05", force=False)
    assert result["status"] == "PARTIAL"
    assert any("schedule workbook has no dates" in m for m in result["missing"])


def test_locked_output_writes_timestamped_copy(_isolate_daily_paths, monkeypatch):
    paths = _isolate_daily_paths
    _make_full_month(paths["drop"], "2026-05")
    out_dir = paths["reports"] / "2026-05"
    out_dir.mkdir(parents=True)
    target = out_dir / "Attendance_Output_Report_2026-05.xlsx"

    real_replace = Path.replace

    def _fake_replace(self, target_path):
        if str(target_path) == str(target):
            raise PermissionError("locked by Excel")
        return real_replace(self, target_path)

    monkeypatch.setattr(Path, "replace", _fake_replace)
    result = daily.run_month("2026-05", force=True)
    daily._write_trial_workbook(target, "2026-05", result)
    assert not target.exists()
    stamped = list(out_dir.glob("Attendance_Output_Report_2026-05_*.xlsx"))
    assert len(stamped) == 1


def test_lock_file_prevents_overlapping_run(_isolate_daily_paths):
    from src.run_lock import RunLocked, acquire_lock, release_lock, _LOCK_PATH
    paths = _isolate_daily_paths
    paths["logs"].mkdir(parents=True, exist_ok=True)
    acquire_lock("first")
    with pytest.raises(RunLocked):
        acquire_lock("second")
    release_lock()


def test_stale_lock_file_is_ignored(_isolate_daily_paths):
    import json as _json
    from src.run_lock import acquire_lock, release_lock, _LOCK_PATH
    paths = _isolate_daily_paths
    paths["logs"].mkdir(parents=True, exist_ok=True)
    lock_path = paths["logs"] / "engine.lock"
    lock_path.write_text(_json.dumps({
        "pid": 99999, "started": time.time() - 7200, "caller": "old",
    }), encoding="utf-8")
    acquire_lock("new")  # must not raise
    release_lock()


def test_open_questions_ignores_info_flags(_isolate_daily_paths):
    paths = _isolate_daily_paths
    from src.rules import load_rules
    rules = load_rules(paths["master"] / "rules.toml")

    info_flag = DayFlag(
        severity="Info", flag="Stayed after scheduled end", person="FAKEONE", home_outlet_tab="AA",
        date=datetime.date(2026, 5, 4), details="", remark="",
        remark_type="Pending question", ref="Q2",
    )
    check_flag = DayFlag(
        severity="Check", flag="Late in (over 15 min)", person="FAKEONE", home_outlet_tab="AA",
        date=datetime.date(2026, 5, 4), details="", remark="",
        remark_type="Pending question", ref="Q1",
    )
    confirmed_flag = DayFlag(
        severity="Check", flag="HQ recorded a later start than the rule allows", person="FAKEONE",
        home_outlet_tab="AA", date=datetime.date(2026, 5, 4), details="", remark="",
        remark_type="Confirmed rule", ref="C6, C5",
    )

    questions = daily._open_questions(rules, [info_flag, check_flag, confirmed_flag])
    by_ref = {q["ref"]: q["flags_affected"] for q in questions}

    assert by_ref.get("Q2", 0) == 0  # the Info flag must not count
    assert by_ref.get("Q1", 0) == 1  # the Check/Pending-question flag counts
    assert "C6" not in by_ref and "C5" not in by_ref  # confirmed-rule refs are not open questions


# ----------------------------------------------------------- rollup (synthetic)

def _day_flag(flag, person, day, details="", remark_type="Possible error"):
    return DayFlag(
        severity="Check", flag=flag, person=person, home_outlet_tab="AA",
        date=datetime.date(2026, 5, day), details=details,
        remark="r", remark_type=remark_type, ref="",
    )


def test_rollup_outlet_flags_one_per_person_with_pairs_and_day_counts():
    outlet = "Outlet differs between schedule and record tab"
    flags = [
        _day_flag(outlet, "FAKEONE", 3, "Schedule block 'AA'; record tab outlet 'BB'."),
        _day_flag(outlet, "FAKEONE", 7, "Schedule block 'AA'; record tab outlet 'BB'."),
        _day_flag(outlet, "FAKEONE", 9, "Schedule block 'CC'; record tab outlet 'BB'."),
        _day_flag(outlet, "FAKETWO", 4, "Schedule block 'AA'; record tab outlet 'AA'."),
    ]
    rolled = daily.rollup_flags(flags)
    assert len(rolled) == 2
    one = rolled[0]
    assert one.person == "FAKEONE" and one.date == "Month" and one.days == 3
    assert "'AA' vs record tab 'BB': 2 day(s), 2026-05-03 to 2026-05-07" in one.details
    assert "'CC' vs record tab 'BB': 1 day(s), 2026-05-09 to 2026-05-09" in one.details
    assert rolled[1].person == "FAKETWO" and rolled[1].days == 1
    assert sum(f.days for f in rolled if f.flag == outlet) == len(flags)


def test_rollup_missing_punch_flags_keep_day_total_and_span():
    flags = [
        _day_flag("Missing clock-in", "FAKEONE", 12),
        _day_flag("Missing clock-in", "FAKEONE", 2),
        _day_flag("Missing clock-out", "FAKEONE", 5),
    ]
    rolled = daily.rollup_flags(flags)
    by_name = {f.flag: f for f in rolled}
    assert by_name["Missing clock-in"].days == 2
    assert by_name["Missing clock-in"].details == "2 day(s), 2026-05-02 to 2026-05-12."
    assert by_name["Missing clock-out"].days == 1
    assert sum(f.days for f in rolled) == 3


def test_rollup_passes_other_flags_through_in_order():
    other = _day_flag("Late in (over 15 min)", "FAKEONE", 6, "x", remark_type="Pending question")
    flags = [
        other,
        _day_flag("Missing clock-in", "FAKEONE", 1),
        _day_flag("Missing clock-in", "FAKEONE", 8),
        _day_flag("Late in (over 15 min)", "FAKEONE", 9, "y", remark_type="Pending question"),
    ]
    rolled = daily.rollup_flags(flags)
    assert [f.flag for f in rolled] == ["Late in (over 15 min)", "Missing clock-in", "Late in (over 15 min)"]
    assert rolled[0] is other and rolled[0].days == 1


def test_rollup_unreadable_outlet_detail_stops_instead_of_guessing():
    flags = [_day_flag("Outlet differs between schedule and record tab", "FAKEONE", 3, "garbled")]
    with pytest.raises(ValueError):
        daily.rollup_flags(flags)


def test_ti_included_with_info_remark_when_exclusion_list_empty():
    staff_rows = {"STAFF_N": {"home_outlet_tab": "AA"}, "FAKEONE": {"home_outlet_tab": "AA"}}
    flags = daily.special_remark_flags(staff_rows, excluded_ids=set())
    assert len(flags) == 1
    assert flags[0].person == "STAFF_N"
    assert flags[0].severity == "Info"
    assert flags[0].date == "Month"
    assert flags[0].remark_type != "Confirmed rule"


def test_ti_remark_absent_when_ti_is_excluded_and_other_staff_never_get_it():
    staff_rows = {"STAFF_N": {"home_outlet_tab": "AA"}, "FAKEONE": {"home_outlet_tab": "AA"}}
    assert daily.special_remark_flags(staff_rows, excluded_ids={"STAFF_N"}) == []
    assert all(f.person == "STAFF_N" for f in daily.special_remark_flags(staff_rows, excluded_ids=set()))


def test_rollup_missing_clock_in_and_out_one_flag_per_person_with_day_span():
    flags = [
        _day_flag("Missing clock-in and clock-out", "FAKEONE", 20),
        _day_flag("Missing clock-in and clock-out", "FAKEONE", 4),
        _day_flag("Missing clock-in and clock-out", "FAKEONE", 11),
        _day_flag("Missing clock-in and clock-out", "FAKETWO", 6),
    ]
    rolled = daily.rollup_flags(flags)
    assert len(rolled) == 2
    one = rolled[0]
    assert one.person == "FAKEONE" and one.severity == "Check"
    assert one.remark_type == "Possible error" and one.date == "Month"
    assert one.days == 3
    assert one.details == "3 day(s), 2026-05-04 to 2026-05-20."
    assert rolled[1].person == "FAKETWO" and rolled[1].days == 1
    assert sum(f.days for f in rolled if f.flag == "Missing clock-in and clock-out") == len(flags)


# ------------------------------------------------------- temporary review page

def _review_path(paths):
    return paths["reports"] / "LATEST" / "Review_May-26.html"


def test_review_page_created_in_latest_with_counts_and_no_external_urls(_isolate_daily_paths):
    paths = _isolate_daily_paths
    _make_full_month(paths["drop"], "2026-05")
    daily.process_month("2026-05", force=True)
    page = _review_path(paths)
    assert page.exists()
    text = page.read_text(encoding="utf-8")
    summary = json.loads((paths["reports"] / "2026-05" / "AGGREGATE_SUMMARY.json").read_text(encoding="utf-8"))
    c = summary["counts"]
    assert f'id="n-red">{c["red"]}</b>Red' in text
    assert f'id="n-amber">{c["amber"]}</b>Amber' in text
    assert f'id="n-green">{c["green"]}</b>Green' in text
    assert f'<b>{summary["info_total"]}</b>Info' in text
    assert "http" not in text.lower()
    assert "<script src" not in text.lower() and "<link" not in text.lower()


def test_review_page_not_copied_into_month_folder(_isolate_daily_paths):
    paths = _isolate_daily_paths
    _make_full_month(paths["drop"], "2026-05")
    daily.process_month("2026-05", force=True)
    assert not list((paths["reports"] / "2026-05").glob("Review_*.html"))


def test_review_page_has_no_staff_name_and_uses_ids_only(_isolate_daily_paths):
    paths = _isolate_daily_paths
    _make_full_month(paths["drop"], "2026-05")
    daily.process_month("2026-05", force=True)
    text = _review_path(paths).read_text(encoding="utf-8")
    assert "Fake One" not in text


def test_review_page_replaced_on_second_run(_isolate_daily_paths):
    paths = _isolate_daily_paths
    _make_full_month(paths["drop"], "2026-05")
    daily.process_month("2026-05", force=True)
    _review_path(paths).write_text("STALE PAGE", encoding="utf-8")
    daily.process_month("2026-05", force=True)
    text = _review_path(paths).read_text(encoding="utf-8")
    assert "STALE PAGE" not in text
    assert "Attendance review" in text


def test_review_hides_name_bearing_details_and_masks_clock_name():
    flag_name = daily.DayFlag(
        severity="Check", flag="Tab name differs from the name typed inside it", person="FAKEONE",
        home_outlet_tab="AA", date="Month", details="expected 'Fake One' but tab reads 'Other Person'.",
        remark="The tab may be copied.", remark_type="Possible error", ref="",
    )
    flag_clock = daily.DayFlag(
        severity="Check", flag="Clock name issue: unknown", person="Some Clock Name", home_outlet_tab="",
        date=datetime.date(2026, 5, 3), details="Clock name 'Some Clock Name' is not in staff_aliases.csv",
        remark="Clock name could not be matched to a staff member.", remark_type="Possible error", ref="",
    )
    summary = {"status": "OK", "counts": {"red": 2, "amber": 0, "green": 0}, "info_total": 0,
               "missing": [], "open_questions": []}
    html_text = daily._review_html("2026-05", summary, [flag_name, flag_clock])
    assert "Fake One" not in html_text and "Other Person" not in html_text
    assert "Some Clock Name" not in html_text
    assert "(unmatched clock name)" in html_text


# --------------------------------------------- pending answers from Client

def _pending_flag(ref, remark_type="Pending question", person="FAKEONE"):
    return daily.DayFlag(
        severity="Check", flag="Late in (over 15 min)", person=person, home_outlet_tab="AA",
        date=datetime.date(2026, 5, 4), details="d", remark="r",
        remark_type=remark_type, ref=ref,
    )


def test_pending_section_present_when_an_open_question_has_flags():
    summary = {"status": "OK", "counts": {"red": 0, "amber": 1, "green": 0}, "info_total": 0,
               "missing": [], "open_questions": [{"ref": "Q1", "rule_ids": ["late_deduction_mode"], "flags_affected": 1}]}
    catalogue = daily.load_pending_catalogue(daily.PENDING_QUESTIONS_PATH, _real_rules())
    pending = daily._pending_questions_for_review(summary, catalogue)
    html_text = daily._review_html("2026-05", summary, [_pending_flag("Q1")], pending)
    assert 'id="pend"' in html_text
    assert "PENDING ANSWERS FROM CLIENT" in html_text


def test_pending_section_absent_when_no_open_question_has_flags():
    summary = {"status": "OK", "counts": {"red": 0, "amber": 0, "green": 0}, "info_total": 0,
               "missing": [], "open_questions": [{"ref": "Q1", "rule_ids": ["late_deduction_mode"], "flags_affected": 0}]}
    catalogue = daily.load_pending_catalogue(daily.PENDING_QUESTIONS_PATH, _real_rules())
    pending = daily._pending_questions_for_review(summary, catalogue)
    html_text = daily._review_html("2026-05", summary, [], pending)
    assert 'id="pend"' not in html_text


def test_pending_rows_carry_data_q():
    summary = {"status": "OK", "counts": {"red": 0, "amber": 1, "green": 0}, "info_total": 0,
               "missing": [], "open_questions": [{"ref": "Q1", "rule_ids": ["late_deduction_mode"], "flags_affected": 1}]}
    html_text = daily._review_html("2026-05", summary, [_pending_flag("Q1")], [])
    assert 'data-q="Q1"' in html_text


def test_pending_page_has_no_network_or_staff_names():
    summary = {"status": "OK", "counts": {"red": 0, "amber": 1, "green": 0}, "info_total": 0,
               "missing": [], "open_questions": [{"ref": "Q1", "rule_ids": ["late_deduction_mode"], "flags_affected": 1}]}
    catalogue = daily.load_pending_catalogue(daily.PENDING_QUESTIONS_PATH, _real_rules())
    pending = daily._pending_questions_for_review(summary, catalogue)
    html_text = daily._review_html("2026-05", summary, [_pending_flag("Q1")], pending)
    assert "http" not in html_text.lower()
    assert "Fake One" not in html_text


def _real_rules():
    from src.rules import load_rules
    return load_rules(MASTER_DATA_DIR_REAL / "rules.toml")


MASTER_DATA_DIR_REAL = Path(__file__).resolve().parent.parent / "03_MASTER_DATA"


def test_pending_catalogue_rejects_unknown_rule_key(tmp_path):
    bad_toml = tmp_path / "pending_questions.toml"
    bad_toml.write_text(
        '[[questions]]\n'
        'ref = "Q1"\n'
        'topic = "x"\n'
        'question = "x?"\n'
        '[[questions.options]]\n'
        'label = "a"\n'
        '[questions.options.effect]\n'
        'no_such_rule_key = true\n',
        encoding="utf-8",
    )
    with pytest.raises(Exception):
        daily.load_pending_catalogue(bad_toml, _real_rules())


# ---- council fixes B1, S8, S10 ------------------------------------------------

def _latest_month(paths):
    summary = paths["reports"] / "LATEST" / "AGGREGATE_SUMMARY.json"
    return json.loads(summary.read_text(encoding="utf-8"))["month"]


def test_empty_month_folder_is_not_ready_and_leaves_latest_alone(_isolate_daily_paths):
    paths = _isolate_daily_paths
    _make_full_month(paths["drop"], "2026-05")
    daily.process_month("2026-05")
    assert _latest_month(paths) == "2026-05"
    (paths["drop"] / "2026-10").mkdir()
    result = daily.process_month("2026-10")
    assert result["status"] == "NOT_READY"
    assert _latest_month(paths) == "2026-05"
    assert not (paths["reports"] / "2026-10").exists()
    assert "2026-10" not in json.loads(daily.MANIFEST_PATH.read_text(encoding="utf-8"))


def test_no_argument_run_with_one_empty_month_keeps_the_good_latest(_isolate_daily_paths):
    paths = _isolate_daily_paths
    _make_full_month(paths["drop"], "2026-05")
    (paths["drop"] / "2026-10").mkdir()
    assert daily.main([]) == 0
    assert _latest_month(paths) == "2026-05"
    last = json.loads((paths["logs"] / "last_run.json").read_text(encoding="utf-8"))
    assert last["statuses"]["2026-10"] == "NOT_READY" and last["latest_updated"] is True


def test_partial_month_does_not_replace_latest_of_another_month(_isolate_daily_paths):
    paths = _isolate_daily_paths
    _make_full_month(paths["drop"], "2026-05")
    daily.process_month("2026-05")
    other = paths["drop"] / "2026-06"
    other.mkdir()
    _write_schedule_workbook(other / "schedule.xlsx", datetime.datetime(2026, 6, 1))
    result = daily.process_month("2026-06")
    assert result["status"] == "PARTIAL" and result["latest_updated"] is False
    assert (paths["reports"] / "2026-06" / "AGGREGATE_SUMMARY.json").exists()
    assert _latest_month(paths) == "2026-05"
    # the same month may refresh its own LATEST, and an OK month always may
    assert daily._may_replace_latest("2026-05", "PARTIAL")
    assert daily._may_replace_latest("2026-06", "OK")
    assert not daily._may_replace_latest("2026-06", "PARTIAL")


def test_nothing_new_run_reports_latest_not_updated(_isolate_daily_paths):
    paths = _isolate_daily_paths
    _make_full_month(paths["drop"], "2026-05")
    daily.main([])
    daily.main([])
    last = json.loads((paths["logs"] / "last_run.json").read_text(encoding="utf-8"))
    assert last["statuses"] == {"2026-05": "NOTHING_NEW"} and last["latest_updated"] is False


def test_run_daily_bat_skips_model_on_last_run_flag():
    text = (Path(__file__).resolve().parents[1] / "run_daily.bat").read_text(encoding="utf-8")
    assert "last_run.json" in text and "latest_updated" in text
    assert "NOTHING_NEW" not in text.split("Step 2", 1)[1]


def test_summary_missing_text_has_no_file_names(_isolate_daily_paths, monkeypatch):
    paths = _isolate_daily_paths
    month_dir = _make_full_month(paths["drop"], "2026-05")
    other_dir = paths["drop"] / "elsewhere"
    other_dir.mkdir()
    duplicate = other_dir / "hr.xlsx"
    duplicate.write_bytes((month_dir / "hr.xlsx").read_bytes())
    monkeypatch.setattr(
        daily, "_classify_month_folder",
        lambda md: ([duplicate], month_dir / "schedule.xlsx", month_dir / "hr.xlsx", []),
    )
    monkeypatch.setattr(daily, "capture_folder", lambda folder, pattern="*": ([], []))
    daily.process_month("2026-05", force=True)
    summary = (paths["reports"] / "2026-05" / "AGGREGATE_SUMMARY.json").read_text(encoding="utf-8")
    assert "duplicate file name" in summary and "hr.xlsx" not in summary


def test_needs_attention_cleared_when_clean_rerun(_isolate_daily_paths):
    paths = _isolate_daily_paths
    month_dir = _make_full_month(paths["drop"], "2026-05")
    wb = openpyxl.Workbook()
    wb.active["A1"] = "nothing useful here"
    wb.save(month_dir / "mystery.xlsx")
    daily.process_month("2026-05")
    note = paths["reports"] / "2026-05" / "NEEDS_ATTENTION.txt"
    assert note.exists()
    (month_dir / "mystery.xlsx").unlink()  # test fixture file, not a real input
    daily.process_month("2026-05", force=True)
    assert not note.exists()
    assert not (paths["reports"] / "LATEST" / "NEEDS_ATTENTION.txt").exists()
