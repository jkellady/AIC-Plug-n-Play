"""Reply handler agent.

Handles a client reply workbook named ``<stem>_reply.xlsx``. The reply
convention (Q4) is unanswered by the client, so it is ASSUMED and the handler
is DISABLED by default: ``run(path)`` with ``enabled=False`` does nothing.

When enabled it parses the stem into a scope and window, finds the original
report under ``06_REPORTS/<YYYY-MM>/``, re-runs the matching reporter, saves the
regenerated workbook next to the original as ``_v2``, ``_v3`` ... (never
overwriting), writes a COUNT-ONLY diff summary to
``07_LOGS/reply_<stem>_<timestamp>.json`` and appends a line to
``07_LOGS/reply_log.csv``.

The reply file is never opened, moved, edited or deleted. The reporters never
overwrite an existing report, so the re-run is pointed at a scratch folder
inside ``06_REPORTS`` (``daily.REPORTS_DIR`` is swapped for the call, which also
keeps the real ``LATEST`` untouched) and the result is copied to the revision
name. The scratch folder is removed afterwards.

Statuses: OK, DISABLED, BLOCKED, NO_MATCH, FAILED. ``run()`` prints nothing.
"""
from __future__ import annotations

import datetime
import importlib
import io
import json
import re
import shutil
from pathlib import Path

import openpyxl

import src.daily as daily

LOG_NAME = "reply_log.csv"
REPLY_SUFFIX = "_reply"

_PATTERNS = [
    ("week", re.compile(r"^(\d{4})-W(\d{2})_Weekly$")),
    ("week", re.compile(r"^Weekly_(\d{4})-W(\d{2})$")),
    ("month", re.compile(r"^(\d{4})-(\d{2})_Monthly$")),
    ("month", re.compile(r"^Monthly_(\d{4})-(\d{2})$")),
    ("day", re.compile(r"^(\d{4})-(\d{2})-(\d{2})_Daily$")),
]


def _parse_stem(stem: str):
    """Return (scope, start, end, canonical_stem, on_disk_stem) or None."""
    for scope, pat in _PATTERNS:
        m = pat.match(stem)
        if not m:
            continue
        try:
            if scope == "week":
                y, w = int(m.group(1)), int(m.group(2))
                start = datetime.date.fromisocalendar(y, w, 1)
                end = start + datetime.timedelta(days=6)
                return scope, start, end, f"{y}-W{w:02d}_Weekly", f"Weekly_{y}-W{w:02d}"
            if scope == "month":
                y, mo = int(m.group(1)), int(m.group(2))
                start = datetime.date(y, mo, 1)
                nxt = datetime.date(y + (mo == 12), mo % 12 + 1, 1)
                end = nxt - datetime.timedelta(days=1)
                return scope, start, end, f"{y}-{mo:02d}_Monthly", f"Monthly_{y}-{mo:02d}"
            d = datetime.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
            return scope, d, d, f"{d.isoformat()}_Daily", f"{d.isoformat()}_Daily"
        except ValueError:
            return None
    return None


def _find_original(reports: Path, end: datetime.date, names: list[str]) -> Path | None:
    home = reports / f"{end.year}-{end.month:02d}"
    for name in names:
        p = home / f"{name}.xlsx"
        if p.is_file():
            return p
    if reports.is_dir():
        for name in names:
            for p in sorted(reports.glob(f"*/{name}.xlsx")):
                if p.is_file():
                    return p
    return None


def _day_rows(data: bytes) -> tuple[dict, bool]:
    """Map (person, date) -> (values..., flag_count) from the Day check sheet."""
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        if "Day check" not in wb.sheetnames:
            return {}, False
        rows = list(wb["Day check"].iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return {}, False
    header = [str(h) for h in rows[0]]
    if "person" not in header or "date" not in header:
        return {}, False
    ip, idt = header.index("person"), header.index("date")
    out = {}
    for r in rows[1:]:
        if r is None or all(v is None for v in r):
            continue
        out[(str(r[ip]), str(r[idt]))] = tuple(str(v) for v in r)
    return out, True


def _flagged(rows: dict) -> int:
    # flag_count is the last column of the Day check sheet
    n = 0
    for vals in rows.values():
        try:
            n += 1 if float(vals[-1]) > 0 else 0
        except (ValueError, IndexError):
            pass
    return n


def _diff(before: bytes, after: bytes) -> dict:
    a, ok_a = _day_rows(before)
    b, ok_b = _day_rows(after)
    added = len(b.keys() - a.keys())
    removed = len(a.keys() - b.keys())
    changed = sum(1 for k in a.keys() & b.keys() if a[k] != b[k])
    return {
        "rows_added": added, "rows_removed": removed, "rows_changed": changed,
        "flagged_person_days_before": _flagged(a), "flagged_person_days_after": _flagged(b),
        "comparable": ok_a and ok_b,
    }


def _write_new(path: Path, data: bytes) -> bool:
    try:
        with open(path, "xb") as f:
            f.write(data)
        return True
    except FileExistsError:
        return False


def _write_revision(original: Path, data: bytes) -> Path:
    n = 2
    while True:
        target = original.with_name(f"{original.stem}_v{n}{original.suffix}")
        if _write_new(target, data):
            return target
        n += 1


def _log(stem: str, status: str, detail: str) -> None:
    logs = Path(daily.LOGS_DIR)
    logs.mkdir(parents=True, exist_ok=True)
    path = logs / LOG_NAME
    is_new = not path.exists()
    detail = detail.replace('"', "'").replace("\n", " ")
    with open(path, "a", encoding="utf-8", newline="") as f:
        if is_new:
            f.write("run_time,stem,status,detail\n")
        f.write(f'{datetime.datetime.now().isoformat(timespec="seconds")},{stem},{status},"{detail}"\n')


def _load_reporter(scope: str):
    mod = "src.agents.weekly_report" if scope == "week" else "src.agents.monthly_report"
    try:
        return importlib.import_module(mod).run, ""
    except Exception as exc:  # missing module or broken import
        return None, f"The {scope}ly reporter is not available yet ({type(exc).__name__}), so no revision was made."


def _call_reporter(fn, scope: str, start: datetime.date) -> dict:
    try:
        return fn(as_of=start, previous=False)
    except TypeError:
        return fn(as_of=start)


def run(reply_path, enabled: bool = False) -> dict:
    result = {"status": "FAILED", "scope": None, "window": None, "revision_path": None,
              "diff_counts": None, "error": ""}
    if not enabled:
        result["status"] = "DISABLED"
        result["error"] = "Reply handling is disabled (reply convention Q4 is unanswered)."
        return result

    name = Path(reply_path).name
    low = name.lower()
    if not low.endswith(REPLY_SUFFIX + ".xlsx"):
        result["error"] = "The file name must look like <report>_reply.xlsx."
        _log("unparsed", "FAILED", result["error"])
        return result
    parsed = _parse_stem(name[: -len(REPLY_SUFFIX + ".xlsx")])
    if parsed is None:
        result["error"] = "The report name in the file name is not a known daily, weekly or monthly report."
        _log("unparsed", "FAILED", result["error"])
        return result
    scope, start, end, stem, disk_stem = parsed
    result["scope"] = scope
    result["window"] = f"{start.isoformat()} to {end.isoformat()}"

    def finish(status: str, error: str = "") -> dict:
        result["status"], result["error"] = status, error
        rev = result["revision_path"]
        _log(stem, status, error or (f"revision={Path(rev).name}" if rev else ""))
        return result

    if not Path(reply_path).is_file():
        return finish("FAILED", "The reply file was not found.")

    reports = Path(daily.REPORTS_DIR)
    original = _find_original(reports, end, [disk_stem, stem])
    if original is None:
        return finish("NO_MATCH", "No original report was found for this reply.")

    if scope == "day":
        return finish("BLOCKED", "There is no daily workbook reporter, so no revision was made.")
    fn, err = _load_reporter(scope)
    if fn is None:
        return finish("BLOCKED", err)

    scratch = reports / f"_reply_tmp_{datetime.datetime.now():%Y%m%d%H%M%S%f}"
    real_reports = daily.REPORTS_DIR
    try:
        scratch.mkdir(parents=True)
        daily.REPORTS_DIR = scratch
        try:
            res = _call_reporter(fn, scope, start)
        finally:
            daily.REPORTS_DIR = real_reports
        out = Path(res.get("output_path") or "")
        if res.get("status") not in ("OK", "PARTIAL") or not out.is_file():
            return finish("FAILED", f"The reporter did not produce a workbook (status {res.get('status')}).")
        new_data = out.read_bytes()
        old_data = original.read_bytes()
        counts = _diff(old_data, new_data)
        revision = _write_revision(original, new_data)
        result["revision_path"] = str(revision)
        result["diff_counts"] = counts

        logs = Path(daily.LOGS_DIR)
        logs.mkdir(parents=True, exist_ok=True)
        stamp = datetime.datetime.now().strftime("%Y%m%dT%H%M%S%f")
        summary = {"stem": stem, "scope": scope, "window": result["window"],
                   "revision_file": revision.name, "diff_counts": counts,
                   "created": datetime.datetime.now().isoformat(timespec="seconds")}
        _write_new(logs / f"reply_{stem}_{stamp}.json", json.dumps(summary, indent=2).encode("utf-8"))
        return finish("OK")
    except Exception as exc:
        return finish("FAILED", f"{type(exc).__name__}: {exc}")
    finally:
        daily.REPORTS_DIR = real_reports
        shutil.rmtree(scratch, ignore_errors=True)


_SUCCESS = {"OK", "DISABLED"}


def main() -> int:
    """Entry point for Task Scheduler. Returns 0 on success, 1 on failure."""
    import sys
    path = sys.argv[1] if len(sys.argv) > 1 else ""
    result = run(path, enabled=bool(path))
    return 0 if result["status"] in _SUCCESS else 1


if __name__ == "__main__":
    import sys
    sys.exit(main())
