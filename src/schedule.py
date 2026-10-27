"""Read the management schedule roster workbook into a long table.

Read-only on 04_REFERENCE_SOURCES: the main pass opens with openpyxl
read_only=True, data_only=True (cached formula values for the hours
column) and streams rows with iter_rows() rather than per-cell lookups;
it never calls save() or writes anything back. A second, non-read_only
handle is opened only to read the small row-1/row-2 header area (merge
ranges for outlet blocks, names) — read_only worksheets don't expose
merged_cells. Nothing is ever written or saved through that handle
either.

Layout (CLAUDE.md section 4): row 1 has outlet block headers (merged
cells), row 2 has staff names with a hours column beside each, data
starts at row 4. Weekly total rows and the "BRING FORWARD HRS" row have
no date in column A and are skipped.
"""
from __future__ import annotations

import datetime
from dataclasses import dataclass
from pathlib import Path

import openpyxl

from src.masterdata import Flag, MasterData, ScheduleCodeRecord

HEADER_ROW = 1
NAME_ROW = 2
FIRST_DATA_ROW = 4

SCHEDULE_LONG_COLUMNS = [
    "person", "date", "outlet_block", "raw_code", "hours", "resolved",
]


@dataclass
class _PersonColumn:
    person: str
    outlet_block: str
    code_col: int
    hours_col: int


def _header_metadata(xlsx_path: Path) -> tuple[dict[int, str], dict[int, str]]:
    """Read row-1 outlet blocks (merge-aware) and row-2 names.

    Opens a small, non-read_only handle just for this (header area only,
    rows 1-2); nothing is ever written or saved through it.
    """
    wb = openpyxl.load_workbook(xlsx_path, read_only=False, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        block_by_col: dict[int, str] = {}
        for merge in ws.merged_cells.ranges:
            if merge.min_row == HEADER_ROW:
                value = ws.cell(HEADER_ROW, merge.min_col).value
                for col in range(merge.min_col, merge.max_col + 1):
                    block_by_col[col] = str(value).strip() if value else ""
        name_by_col: dict[int, str] = {}
        for col in range(1, ws.max_column + 1):
            value = ws.cell(NAME_ROW, col).value
            if value:
                name_by_col[col] = str(value).strip()
        return block_by_col, name_by_col
    finally:
        wb.close()


def _person_columns(block_by_col: dict[int, str], name_by_col: dict[int, str]) -> list[_PersonColumn]:
    columns: list[_PersonColumn] = []
    for col, name in sorted(name_by_col.items()):
        if col < 3 or name.upper() == "REMARK":
            continue
        columns.append(_PersonColumn(
            person=name,
            outlet_block=block_by_col.get(col, ""),
            code_col=col,
            hours_col=col + 1,
        ))
    return columns


def _parse_date(value):
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    return None


def parse_schedule_workbook(
    xlsx_path: Path, master_data: MasterData
) -> tuple[list[dict], int]:
    """Parse the schedule workbook into a long table.

    Returns (rows, skipped_row_count). Each row is a dict with keys
    SCHEDULE_LONG_COLUMNS; "resolved" is a ScheduleCodeRecord or a Flag
    (unknown codes, including asterisked/arrow codes not in
    schedule_codes.csv, come out as Flags with the raw code kept).
    """
    block_by_col, name_by_col = _header_metadata(xlsx_path)
    columns = _person_columns(block_by_col, name_by_col)

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows: list[dict] = []
        skipped = 0
        for row_values in ws.iter_rows(min_row=FIRST_DATA_ROW, values_only=True):
            date = _parse_date(row_values[0] if row_values else None)
            if date is None:
                skipped += 1
                continue
            for pc in columns:
                raw_code = row_values[pc.code_col - 1] if pc.code_col - 1 < len(row_values) else None
                if raw_code is None or str(raw_code).strip() == "":
                    continue
                raw_code = str(raw_code).strip()
                hours = row_values[pc.hours_col - 1] if pc.hours_col - 1 < len(row_values) else None
                resolved = master_data.lookup_schedule_code(raw_code)
                rows.append({
                    "person": pc.person,
                    "date": date,
                    "outlet_block": pc.outlet_block,
                    "raw_code": raw_code,
                    "hours": hours,
                    "resolved": resolved,
                })
        return rows, skipped
    finally:
        wb.close()
