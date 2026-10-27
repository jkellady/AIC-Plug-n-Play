"""Read-only reader for the HR record workbook (Time Calculation tabs).

Read-only on 04_REFERENCE_SOURCES: both passes open with openpyxl
read_only=True and stream rows with iter_rows() (no per-cell ws.cell()
lookups on a normal-mode workbook, which is what made the schedule
reader slow before it was fixed). Never calls save().

Layout, confirmed by inspecting the real May 2026 workbook (25 sheets:
18 real person tabs for May 2026, plus `xxx`, `PT_OUT3` and five `DRAFT_*`
tabs that are stale templates for other months):
  - row 2 col D: the tab's own month (a datetime). Only tabs whose month
    matches the target (year, month) are parsed — row 8 col D == "IN"
    alone does NOT distinguish real tabs from drafts, since every one of
    the 25 sheets has "IN" there.
  - row 3 col D: the name typed in the tab (name_in_tab). row 3 col H:
    the branch/outlet text (tab_outlet).
  - row 8: header row. D=IN, E=OUT, F=Working(Hr:Min), H=OT 1.5,
    O..R = four raw clock-punch columns, in order.
  - rows 9..39: one row per day, day number in col C (day = row - 8).
    Rows for a day number beyond the target month's real day count are
    skipped (handles a 30-day month's unused day-31 row).

D (IN), E (OUT), F (Working) and H (OT 1.5) are each either a typed
value or a formula (e.g. F is a typed "08:00:00" instead of the usual
"=E-D" on some leave days — this is what a step 4d flag like "day typed
8:00 instead of 8:30" needs to tell apart). This reader records both the
cached value and whether the cell held a formula, for each of the four.
"""
from __future__ import annotations

import calendar
import datetime
from dataclasses import dataclass

import openpyxl

HEADER_IN_ROW = 8
HEADER_IN_COL = 4  # D
NAME_ROW = 3
NAME_COL = 4  # D
OUTLET_LABEL_COL = 7  # G ("Branch -")
OUTLET_VALUE_COL = 8  # H
MONTH_ROW = 2
MONTH_COL = 4  # D
FIRST_DATA_ROW = 9
DAY_COL = 3  # C

COL_IN = 4       # D
COL_OUT = 5      # E
COL_WORKING = 6  # F
COL_OT_1_5 = 8   # H
PUNCH_COLS = (15, 16, 17, 18)  # O, P, Q, R

HR_ROW_COLUMNS = [
    "person", "name_in_tab", "tab_outlet", "date",
    "hq_in", "hq_in_raw", "hq_in_is_formula",
    "hq_out", "hq_out_raw", "hq_out_is_formula",
    "hq_working_hours", "hq_working_hours_raw", "hq_working_hours_is_formula",
    "hq_ot_1_5", "hq_ot_1_5_raw", "hq_ot_1_5_is_formula",
    "raw_punches",
]


@dataclass
class _SheetHeader:
    name_in_tab: str
    tab_outlet: str
    month_year: int
    month_month: int


def _is_formula(raw_value) -> bool:
    return isinstance(raw_value, str) and raw_value.startswith("=")


def _to_time_str(value) -> str | None:
    if isinstance(value, datetime.datetime):
        value = value.time()
    if isinstance(value, datetime.time):
        return f"{value.hour:02d}:{value.minute:02d}"
    return None


def _to_decimal_hours(value) -> float | None:
    if isinstance(value, datetime.timedelta):
        return round(value.total_seconds() / 3600, 4)
    if isinstance(value, datetime.time):
        return round(value.hour + value.minute / 60 + value.second / 3600, 4)
    return None


def _raw_text(value) -> str:
    return "" if value is None else str(value)


def _read_sheet_rows(ws, values_only: bool = True):
    """Read a whole sheet into a dict keyed by (row, col) -> value, via iter_rows."""
    grid: dict[tuple[int, int], object] = {}
    for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
        for c, value in enumerate(row, start=1):
            if value is not None:
                grid[(r, c)] = value
    return grid


def _sheet_header(grid: dict[tuple[int, int], object]) -> _SheetHeader | None:
    month_value = grid.get((MONTH_ROW, MONTH_COL))
    if not isinstance(month_value, (datetime.date, datetime.datetime)):
        return None
    return _SheetHeader(
        name_in_tab=str(grid.get((NAME_ROW, NAME_COL)) or ""),
        tab_outlet=str(grid.get((NAME_ROW, OUTLET_VALUE_COL)) or ""),
        month_year=month_value.year,
        month_month=month_value.month,
    )


def read_hr_workbook(
    xlsx_path, year: int, month: int
) -> tuple[list[dict], list[tuple[str, str]]]:
    """Read every person tab for (year, month) into a long table.

    Returns (rows, skipped) where skipped is a list of (sheet_name, reason)
    for every sheet that was not parsed as a person tab for this month.
    """
    values_wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    formulas_wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=False)
    try:
        days_in_month = calendar.monthrange(year, month)[1]
        rows: list[dict] = []
        skipped: list[tuple[str, str]] = []

        for sheet_name in values_wb.sheetnames:
            values_grid = _read_sheet_rows(values_wb[sheet_name])
            header = _sheet_header(values_grid)
            if header is None:
                skipped.append((sheet_name, "no date in row 2 col D"))
                continue
            if (header.month_year, header.month_month) != (year, month):
                skipped.append((
                    sheet_name,
                    f"row 2 month is {header.month_year}-{header.month_month:02d}, "
                    f"not {year}-{month:02d}",
                ))
                continue
            if values_grid.get((HEADER_IN_ROW, HEADER_IN_COL)) != "IN":
                skipped.append((sheet_name, "row 8 col D is not 'IN'"))
                continue

            formulas_grid = _read_sheet_rows(formulas_wb[sheet_name])

            for day in range(1, days_in_month + 1):
                row_num = FIRST_DATA_ROW + (day - 1)
                day_cell = values_grid.get((row_num, DAY_COL))
                if day_cell != day:
                    continue  # stray or missing day-number cell; nothing to emit

                in_val = values_grid.get((row_num, COL_IN))
                out_val = values_grid.get((row_num, COL_OUT))
                working_val = values_grid.get((row_num, COL_WORKING))
                ot15_val = values_grid.get((row_num, COL_OT_1_5))

                raw_punches = [
                    _to_time_str(values_grid.get((row_num, c))) for c in PUNCH_COLS
                ]

                rows.append({
                    "person": sheet_name,
                    "name_in_tab": header.name_in_tab,
                    "tab_outlet": header.tab_outlet,
                    "date": datetime.date(year, month, day),
                    "hq_in": _to_time_str(in_val),
                    "hq_in_raw": _raw_text(in_val),
                    "hq_in_is_formula": _is_formula(formulas_grid.get((row_num, COL_IN))),
                    "hq_out": _to_time_str(out_val),
                    "hq_out_raw": _raw_text(out_val),
                    "hq_out_is_formula": _is_formula(formulas_grid.get((row_num, COL_OUT))),
                    "hq_working_hours": _to_decimal_hours(working_val),
                    "hq_working_hours_raw": _raw_text(working_val),
                    "hq_working_hours_is_formula": _is_formula(
                        formulas_grid.get((row_num, COL_WORKING))
                    ),
                    "hq_ot_1_5": _to_decimal_hours(ot15_val),
                    "hq_ot_1_5_raw": _raw_text(ot15_val),
                    "hq_ot_1_5_is_formula": _is_formula(formulas_grid.get((row_num, COL_OT_1_5))),
                    "raw_punches": raw_punches,
                })

        return rows, skipped
    finally:
        values_wb.close()
        formulas_wb.close()
